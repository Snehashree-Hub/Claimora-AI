from pathlib import Path
from datetime import datetime

from openpyxl import Workbook, load_workbook


EXCEL_PATH = Path(__file__).resolve().parents[1] / "data" / "claims.xlsx"


EXCEL_HEADERS = [
    "id",
    "claim_number",
    "policy_number",
    "user_id",
    "patient_name",
    "age",
    "insurance_type",
    "insurance_company",
    "hospital_name",
    "treatment",
    "claim_amount",
    "eligibility_status",
    "document_status",
    "fraud_risk",
    "assessment_status",
    "email_status",
    "overall_status",
    "remarks",
    "created_at",
]


def append_claim_to_excel(claim):
    EXCEL_PATH.parent.mkdir(parents=True, exist_ok=True)

    if EXCEL_PATH.exists():
        workbook = load_workbook(EXCEL_PATH)
        worksheet = workbook.active

        if worksheet.max_row == 1 and worksheet.cell(1, 1).value is None:
            for column, header in enumerate(EXCEL_HEADERS, 1):
                worksheet.cell(1, column).value = header

    else:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Claims"

        for column, header in enumerate(EXCEL_HEADERS, 1):
            worksheet.cell(1, column).value = header

    existing_ids = set()

    for row in range(2, worksheet.max_row + 1):
        existing_ids.add(worksheet.cell(row, 1).value)

    if claim.id in existing_ids:
        workbook.close()
        return True

    values = [
        claim.id,
        claim.claim_number,
        claim.policy_number,
        claim.user_id,
        claim.patient_name,
        claim.age,
        claim.insurance_type,
        claim.insurance_company,
        claim.hospital_name,
        claim.treatment,
        claim.claim_amount,
        claim.eligibility_status,
        claim.document_status,
        claim.fraud_risk,
        claim.assessment_status,
        claim.email_status,
        claim.overall_status,
        claim.remarks,
        datetime.utcnow().isoformat(),
    ]

    worksheet.append(values)

    workbook.save(EXCEL_PATH)
    workbook.close()

    return True
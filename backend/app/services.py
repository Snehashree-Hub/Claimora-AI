import os
from datetime import datetime
from openpyxl import Workbook, load_workbook

# ============================================================
# EMAIL CONFIGURATION
# ============================================================

EMAIL_MODE = os.getenv("EMAIL_MODE", "DEMO").upper()

# ============================================================
# EXCEL CONFIGURATION
# ============================================================

EXCEL_FOLDER = "data"
EXCEL_FILE = os.path.join(EXCEL_FOLDER, "claims.xlsx")

# ============================================================
# EXTERNAL INSURANCE COMPANIES
# DEMO ROUTING ONLY
# ============================================================

INSURANCE_COMPANIES = {
    "Star Health": {
        "department": "Star Health Claims Department"
    },
    "Niva Bupa": {
        "department": "Niva Bupa Claims Department"
    },
    "Aditya Birla Health": {
        "department": "Aditya Birla Health Claims Department"
    },
    "HDFC ERGO": {
        "department": "HDFC ERGO Claims Department"
    },
    "ICICI Lombard": {
        "department": "ICICI Lombard Claims Department"
    },
    "Care Health": {
        "department": "Care Health Claims Department"
    },
    "Bajaj Allianz": {
        "department": "Bajaj Allianz Claims Department"
    },
    "SBI General": {
        "department": "SBI General Claims Department"
    }
}


# ============================================================
# CREATE EXCEL FILE
# ============================================================

def create_excel_file():

    os.makedirs(EXCEL_FOLDER, exist_ok=True)

    if not os.path.exists(EXCEL_FILE):

        workbook = Workbook()

        sheet = workbook.active
        sheet.title = "Claims"

        headers = [
            "Claim ID",
            "Patient Name",
            "Age",
            "Policy Number",
            "Insurance Company",
            "Hospital Name",
            "Treatment",
            "Claim Amount",
            "Eligibility",
            "Documents",
            "Fraud Risk",
            "Email Status",
            "Overall Status",
            "Submitted Date"
        ]

        sheet.append(headers)

        workbook.save(EXCEL_FILE)


# ============================================================
# SAVE CLAIM TO EXCEL
# ============================================================

def save_claim_to_excel(claim):

    create_excel_file()

    workbook = load_workbook(EXCEL_FILE)

    sheet = workbook["Claims"]

    sheet.append([
        claim.claim_number,
        claim.patient_name,
        claim.age,
        claim.policy_number,
        claim.insurance_company,
        claim.hospital_name,
        claim.treatment,
        claim.claim_amount,
        claim.eligibility_status,
        claim.document_status,
        claim.fraud_risk,
        claim.email_status,
        claim.overall_status,
        datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    ])

    workbook.save(EXCEL_FILE)


# ============================================================
# CREATE CLAIM EMAIL
# ============================================================

def build_claim_email(claim):

    subject = f"Health Insurance Claim - {claim.claim_number}"

    email_body = f"""
Dear Claims Department,

A new health insurance claim has been submitted
through ClaimCare's AI Insurance Claim Assistance System.

------------------------------------------------------------
CLAIM INFORMATION
------------------------------------------------------------

Claim ID:
{claim.claim_number}

Patient Name:
{claim.patient_name}

Patient Age:
{claim.age}

Policy Number:
{claim.policy_number}

Insurance Company:
{claim.insurance_company}

Hospital:
{claim.hospital_name}

Treatment:
{claim.treatment}

Claim Amount:
Rs. {claim.claim_amount:,.2f}

------------------------------------------------------------
AI VERIFICATION STATUS
------------------------------------------------------------

Eligibility:
{claim.eligibility_status}

Document Verification:
{claim.document_status}

Fraud Risk:
{claim.fraud_risk}

Assessment:
{claim.assessment_status}

Overall Status:
{claim.overall_status}

------------------------------------------------------------

Please review this claim according to the
applicable insurance policy terms and procedures.

Regards,

ClaimCare
AI Insurance Claim Assistance System
"""

    return subject, email_body


# ============================================================
# DEMO EMAIL
# ============================================================

def send_demo_email(claim, recipient_department):

    subject, email_body = build_claim_email(claim)

    print()
    print("=" * 70)
    print("              AI INSURANCE CLAIM EMAIL")
    print("=" * 70)

    print(f"EMAIL MODE       : DEMO")
    print(f"Claim ID         : {claim.claim_number}")
    print(f"Patient          : {claim.patient_name}")
    print(f"Insurance        : {claim.insurance_company}")
    print(f"Recipient        : {recipient_department}")
    print(f"Subject          : {subject}")

    print("-" * 70)
    print("EMAIL BODY")
    print("-" * 70)

    print(email_body)

    print("-" * 70)
    print("EMAIL STATUS     : SIMULATED")
    print("NO REAL EMAIL WAS SENT")
    print("=" * 70)
    print()

    return {
        "success": True,
        "status": "Simulated",
        "mode": "DEMO",
        "recipient": recipient_department,
        "subject": subject,
        "message": "Claim email simulated successfully. No real email was sent."
    }


# ============================================================
# SEND CLAIM EMAIL
# ============================================================

def send_claim_email(claim):

    insurance = INSURANCE_COMPANIES.get(
        claim.insurance_company
    )

    if insurance is None:

        return {
            "success": False,
            "status": "Failed",
            "mode": EMAIL_MODE,
            "message": "Insurance company is not configured."
        }

    recipient_department = insurance["department"]

    # --------------------------------------------------------
    # COLLEGE PROJECT DEMO MODE
    # --------------------------------------------------------

    if EMAIL_MODE == "DEMO":

        return send_demo_email(
            claim,
            recipient_department
        )

    # --------------------------------------------------------
    # REAL EMAIL MODE IS DISABLED FOR THIS COLLEGE DEMO
    # --------------------------------------------------------

    return {
        "success": False,
        "status": "Not Configured",
        "mode": EMAIL_MODE,
        "message": (
            "Real SMTP email is not configured. "
            "Use EMAIL_MODE=DEMO for the college project."
        )
    }

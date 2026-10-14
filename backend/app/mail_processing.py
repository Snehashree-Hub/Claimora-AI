import re
from datetime import datetime
from pathlib import Path
from threading import Lock
from typing import Optional

from fastapi import APIRouter, HTTPException
from pydantic import BaseModel
from openpyxl import Workbook, load_workbook


router = APIRouter(
    prefix="/api/mail",
    tags=["Mail Processing"],
)


# ============================================================
# CONFIGURATION
# ============================================================

BASE_DIR = Path(__file__).resolve().parent.parent
DATA_DIR = BASE_DIR / "data"
DATA_DIR.mkdir(parents=True, exist_ok=True)

MAIL_EXCEL_FILE = DATA_DIR / "mail_processing.xlsx"

_excel_lock = Lock()


# ============================================================
# MODELS
# ============================================================

class MailProcessRequest(BaseModel):
    sender_email: str
    recipient_email: Optional[str] = None
    subject: str
    body: str


class MailResponse(BaseModel):
    id: int
    received_at: str
    sender_email: str
    recipient_email: Optional[str]
    subject: str
    body: str
    mail_type: str
    claim_id: Optional[str]
    policy_id: Optional[str]
    patient_email: Optional[str]
    requested_documents: str
    extracted_information: str
    matched_claim: bool
    processing_status: str
    action_required: str


# ============================================================
# EXCEL
# ============================================================

MAIL_HEADERS = [
    "ID",
    "Received At",
    "Sender Email",
    "Recipient Email",
    "Subject",
    "Body",
    "Mail Type",
    "Claim ID",
    "Policy ID",
    "Patient Email",
    "Requested Documents",
    "Extracted Information",
    "Matched Claim",
    "Processing Status",
    "Action Required",
]


def ensure_mail_workbook():
    with _excel_lock:

        if MAIL_EXCEL_FILE.exists():
            return

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Mail Processing"

        sheet.append(MAIL_HEADERS)

        workbook.save(MAIL_EXCEL_FILE)


def read_mail_rows():
    ensure_mail_workbook()

    with _excel_lock:
        workbook = load_workbook(
            MAIL_EXCEL_FILE,
            read_only=True,
            data_only=True,
        )

        sheet = workbook["Mail Processing"]

        rows = list(sheet.iter_rows(values_only=True))

        workbook.close()

    result = []

    if not rows:
        return result

    headers = list(rows[0])

    for row in rows[1:]:
        if not any(value is not None for value in row):
            continue

        item = {}

        for index, header in enumerate(headers):
            value = row[index] if index < len(row) else None
            item[header] = value

        result.append(item)

    return result


def append_mail_to_excel(mail_data: dict):

    ensure_mail_workbook()

    with _excel_lock:

        workbook = load_workbook(MAIL_EXCEL_FILE)

        sheet = workbook["Mail Processing"]

        next_id = sheet.max_row

        sheet.append([
            next_id,
            mail_data["received_at"],
            mail_data["sender_email"],
            mail_data.get("recipient_email"),
            mail_data["subject"],
            mail_data["body"],
            mail_data["mail_type"],
            mail_data.get("claim_id"),
            mail_data.get("policy_id"),
            mail_data.get("patient_email"),
            mail_data.get("requested_documents"),
            mail_data.get("extracted_information"),
            "Yes" if mail_data.get("matched_claim") else "No",
            mail_data["processing_status"],
            mail_data["action_required"],
        ])

        workbook.save(MAIL_EXCEL_FILE)
        workbook.close()

    return next_id


# ============================================================
# EMAIL CLASSIFICATION
# ============================================================

def classify_email(subject: str, body: str) -> str:

    text = f"{subject} {body}".lower()

    document_words = [
        "document",
        "documents",
        "discharge summary",
        "medical bill",
        "hospital bill",
        "final bill",
        "prescription",
        "invoice",
        "report",
        "upload",
        "submit",
    ]

    status_words = [
        "status",
        "claim status",
        "approved",
        "rejected",
        "under review",
        "processing",
        "pending",
    ]

    information_words = [
        "additional information",
        "more information",
        "information required",
        "clarification",
        "clarification required",
        "additional details",
    ]

    approval_words = [
        "approved",
        "approval",
    ]

    rejection_words = [
        "rejected",
        "rejection",
    ]

    if any(word in text for word in information_words):
        return "Additional Information Request"

    if any(word in text for word in document_words):
        return "Document Request"

    if any(word in text for word in approval_words):
        return "Claim Decision Notification"

    if any(word in text for word in rejection_words):
        return "Claim Decision Notification"

    if any(word in text for word in status_words):
        return "Claim Status Update"

    return "General Claim Communication"


# ============================================================
# CLAIM / POLICY EXTRACTION
# ============================================================

def extract_claim_id(text: str) -> Optional[str]:

    patterns = [
        r"\bCLM[- ]?\d{4}[- ]?\d{1,8}\b",
        r"\bCLM[- ]?\d{1,10}\b",
        r"\bDEMO-CLM-\d{3}\b",
    ]

    for pattern in patterns:

        match = re.search(
            pattern,
            text,
            flags=re.IGNORECASE,
        )

        if match:
            return match.group(0).upper().replace(" ", "")

    return None


def extract_policy_id(text: str) -> Optional[str]:

    patterns = [
        r"\bPOL[- ]?\d{4}[- ]?\d{1,8}\b",
        r"\bPOL[- ]?\d{1,10}\b",
        r"\bDEMO-POL-\d{3}\b",
        r"\bDEMO-RKH-\d{3}\b",
    ]

    for pattern in patterns:

        match = re.search(
            pattern,
            text,
            flags=re.IGNORECASE,
        )

        if match:
            return match.group(0).upper().replace(" ", "")

    return None


def extract_email(text: str) -> Optional[str]:

    match = re.search(
        r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}",
        text,
    )

    if match:
        return match.group(0).lower()

    return None


# ============================================================
# DOCUMENT DETECTION
# ============================================================

def extract_requested_documents(text: str):

    text_lower = text.lower()

    document_mapping = {
        "discharge summary": [
            "discharge summary",
            "discharge report",
        ],
        "final hospital bill": [
            "final bill",
            "final hospital bill",
            "hospital bill",
        ],
        "medical report": [
            "medical report",
            "medical reports",
        ],
        "prescription": [
            "prescription",
            "prescription copy",
        ],
        "invoice": [
            "invoice",
        ],
        "diagnostic report": [
            "diagnostic report",
            "diagnostic reports",
            "lab report",
        ],
        "treatment records": [
            "treatment records",
            "treatment record",
        ],
        "identity proof": [
            "identity proof",
            "id proof",
            "identity document",
        ],
    }

    found = []

    for document_name, keywords in document_mapping.items():

        for keyword in keywords:

            if keyword in text_lower:
                found.append(document_name)
                break

    return sorted(set(found))


# ============================================================
# INFORMATION EXTRACTION
# ============================================================

def build_extracted_information(
    claim_id,
    policy_id,
    patient_email,
    documents,
):

    information = []

    if claim_id:
        information.append(f"Claim ID: {claim_id}")

    if policy_id:
        information.append(f"Policy ID: {policy_id}")

    if patient_email:
        information.append(
            f"Patient Email: {patient_email}"
        )

    if documents:
        information.append(
            "Requested Documents: "
            + ", ".join(documents)
        )

    return " | ".join(information)


# ============================================================
# CLAIM MATCHING
# ============================================================

def check_claim_match(
    claim_id: Optional[str],
    policy_id: Optional[str],
    sender_email: Optional[str],
):

    try:

        from .database import SessionLocal
        from .models import Claim

        db = SessionLocal()

        query = db.query(Claim)

        claim = None

        if claim_id:

            claim = (
                query
                .filter(
                    Claim.claim_id == claim_id
                )
                .first()
            )

        if not claim and policy_id:

            claim = (
                query
                .filter(
                    Claim.policy_id == policy_id
                )
                .first()
            )

        if not claim and sender_email:

            claim = (
                query
                .filter(
                    Claim.owner_email == sender_email
                )
                .order_by(Claim.id.desc())
                .first()
            )

        db.close()

        return claim

    except Exception:

        return None


# ============================================================
# PROCESS MAIL
# ============================================================

@router.post(
    "/process",
    response_model=MailResponse,
)
def process_mail(
    request: MailProcessRequest,
):

    subject = request.subject.strip()
    body = request.body.strip()
    sender = request.sender_email.strip().lower()

    combined_text = f"{subject}\n{body}"

    mail_type = classify_email(
        subject,
        body,
    )

    claim_id = extract_claim_id(
        combined_text
    )

    policy_id = extract_policy_id(
        combined_text
    )

    patient_email = extract_email(
        body
    )

    if not patient_email:
        patient_email = sender

    requested_documents = (
        extract_requested_documents(
            combined_text
        )
    )

    matched_claim = check_claim_match(
        claim_id=claim_id,
        policy_id=policy_id,
        sender_email=patient_email,
    )

    extracted_information = (
        build_extracted_information(
            claim_id=claim_id,
            policy_id=policy_id,
            patient_email=patient_email,
            documents=requested_documents,
        )
    )

    # ========================================================
    # PROCESSING LOGIC
    # ========================================================

    if mail_type == "Document Request":

        processing_status = (
            "Additional Documents Required"
        )

        if requested_documents:

            action_required = (
                "Notify patient to submit: "
                + ", ".join(requested_documents)
            )

        else:

            action_required = (
                "Review requested documents"
            )

    elif mail_type == "Additional Information Request":

        processing_status = (
            "Additional Information Required"
        )

        action_required = (
            "Review email and collect requested information"
        )

    elif mail_type == "Claim Decision Notification":

        processing_status = (
            "Decision Communication Received"
        )

        action_required = (
            "Review authorised insurer decision"
        )

    elif mail_type == "Claim Status Update":

        processing_status = (
            "Claim Status Communication Received"
        )

        action_required = (
            "Review claim status and notify patient if required"
        )

    else:

        processing_status = (
            "Mail Processed"
        )

        action_required = (
            "Review communication"
        )

    received_at = datetime.now().isoformat(
        timespec="seconds"
    )

    mail_data = {
        "received_at": received_at,
        "sender_email": sender,
        "recipient_email": request.recipient_email,
        "subject": subject,
        "body": body,
        "mail_type": mail_type,
        "claim_id": claim_id,
        "policy_id": policy_id,
        "patient_email": patient_email,
        "requested_documents": ", ".join(
            requested_documents
        ),
        "extracted_information": extracted_information,
        "matched_claim": matched_claim is not None,
        "processing_status": processing_status,
        "action_required": action_required,
    }

    mail_id = append_mail_to_excel(
        mail_data
    )

    return {
        "id": mail_id,
        **mail_data,
    }


# ============================================================
# GET MAIL HISTORY
# ============================================================

@router.get("/")
def get_processed_mails():

    rows = read_mail_rows()

    result = []

    for row in reversed(rows):

        result.append({
            "id": row.get("ID"),
            "received_at": row.get("Received At"),
            "sender_email": row.get("Sender Email"),
            "recipient_email": row.get("Recipient Email"),
            "subject": row.get("Subject"),
            "body": row.get("Body"),
            "mail_type": row.get("Mail Type"),
            "claim_id": row.get("Claim ID"),
            "policy_id": row.get("Policy ID"),
            "patient_email": row.get("Patient Email"),
            "requested_documents": row.get("Requested Documents"),
            "extracted_information": row.get(
                "Extracted Information"
            ),
            "matched_claim": str(
                row.get("Matched Claim", "")
            ).lower() == "yes",
            "processing_status": row.get(
                "Processing Status"
            ),
            "action_required": row.get(
                "Action Required"
            ),
        })

    return result


# ============================================================
# MAIL SUMMARY
# ============================================================

@router.get("/summary")
def mail_summary():

    rows = read_mail_rows()

    total = len(rows)

    document_requests = sum(
        1
        for row in rows
        if row.get("Mail Type")
        == "Document Request"
    )

    information_requests = sum(
        1
        for row in rows
        if row.get("Mail Type")
        == "Additional Information Request"
    )

    status_updates = sum(
        1
        for row in rows
        if row.get("Mail Type")
        == "Claim Status Update"
    )

    matched = sum(
        1
        for row in rows
        if str(
            row.get("Matched Claim", "")
        ).lower()
        == "yes"
    )

    return {
        "total": total,
        "document_requests": document_requests,
        "information_requests": information_requests,
        "status_updates": status_updates,
        "matched_claims": matched,
        "excel_file": str(
            MAIL_EXCEL_FILE
        ),
    }


# ============================================================
# TEST EMAIL
# ============================================================

@router.post("/test")
def create_test_mail():

    request = MailProcessRequest(
        sender_email="claims@insurance-demo.com",
        recipient_email="claims@rkhospital.com",
        subject=(
            "Additional documents required "
            "for CLM-2026-000001"
        ),
        body=(
            "Dear Patient,\n\n"
            "Please submit the discharge summary "
            "and final hospital bill for "
            "claim CLM-2026-000001.\n\n"
            "Policy ID: POL-2026-000001\n"
            "Please provide the documents "
            "within the required processing period."
        ),
    )

    return process_mail(request)


# ============================================================
# EXCEL STATUS
# ============================================================

@router.get("/excel/status")
def mail_excel_status():

    ensure_mail_workbook()

    return {
        "status": "ready",
        "file": str(
            MAIL_EXCEL_FILE
        ),
        "exists": MAIL_EXCEL_FILE.exists(),
    }